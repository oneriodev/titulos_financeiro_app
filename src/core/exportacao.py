
"""
Exportação dos dados tratados para CSV e XLSX.

Os arquivos são gerados em memória, sem armazenamento
no servidor. A exportação recebe o DataFrame já filtrado,
preservando os dados numéricos e as datas.
"""

import io

import pandas as pd


def _proteger_textos(df: pd.DataFrame) -> pd.DataFrame:
    """
    Neutraliza possíveis fórmulas em campos textuais.

    Não modifica o DataFrame original. A proteção é aplicada
    somente à cópia destinada à exportação.
    """
    resultado = df.copy()

    for coluna in resultado.select_dtypes(
        include=["object", "string"]
    ).columns:
        def proteger(valor):
            if not isinstance(valor, str):
                return valor

            # Espaços e controles iniciais também podem ocultar fórmulas.
            if valor.lstrip().startswith(("=", "+", "-", "@")):
                return "'" + valor

            return valor

        resultado[coluna] = resultado[coluna].map(proteger)

    return resultado


def exportar_csv(df: pd.DataFrame) -> bytes:
    """
    Gera CSV compatível com o padrão brasileiro.

    O BOM facilita o reconhecimento do UTF-8 pelo Excel.
    O ponto e vírgula evita conflito com a vírgula decimal.
    """
    dados = _proteger_textos(df)

    conteudo = dados.to_csv(
        index=False,
        sep=";",
        decimal=",",
        encoding="utf-8-sig",
        date_format="%d/%m/%Y",
    )

    return conteudo.encode("utf-8-sig")





def exportar_xlsx(df: pd.DataFrame) -> bytes:
    """
    Exporta os dados tratados para XLSX em memória.

    Datas são exportadas sem o componente de horário.
    Valores monetários permanecem numéricos e o DataFrame
    original não é modificado.
    """
    from datetime import date, datetime

    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    # Trabalha com uma cópia protegida dos dados.
    dados = _proteger_textos(df)

    # Remove o horário das colunas datetime somente
    # na cópia destinada à exportação.
    for coluna in dados.columns:
        if pd.api.types.is_datetime64_any_dtype(
            dados[coluna]
        ):
            dados[coluna] = dados[coluna].dt.date

    buffer = io.BytesIO()

    with pd.ExcelWriter(
        buffer,
        engine="openpyxl",
        date_format="DD/MM/YYYY",
        datetime_format="DD/MM/YYYY",
    ) as escritor:

        dados.to_excel(
            escritor,
            sheet_name="Dados Tratados",
            index=False,
        )

        planilha = escritor.sheets["Dados Tratados"]

        planilha.freeze_panes = "A2"
        planilha.auto_filter.ref = planilha.dimensions

        for indice, coluna in enumerate(
            dados.columns, start=1
        ):
            letra = get_column_letter(indice)

            cabecalho = planilha.cell(
                row=1,
                column=indice,
            )
            cabecalho.font = Font(bold=True)

            # Verifica as células efetivamente gravadas,
            # e não apenas o tipo original da coluna.
            coluna_tem_data = False
            maior_comprimento = len(str(coluna))

            for linha in range(
                2, planilha.max_row + 1
            ):
                celula = planilha.cell(
                    row=linha,
                    column=indice,
                )

                if isinstance(
                    celula.value, (date, datetime)
                ):
                    coluna_tem_data = True
                    celula.number_format = "DD/MM/YYYY"
                    celula.alignment = Alignment(
                        horizontal="center"
                    )
                elif celula.value is not None:
                    maior_comprimento = max(
                        maior_comprimento,
                        len(str(celula.value)),
                    )

            if coluna_tem_data:
                # Espaço suficiente para DD/MM/YYYY.
                largura = max(
                    18,
                    len(str(coluna)) + 3,
                )
            else:
                largura = min(
                    max(maior_comprimento + 3, 12),
                    50,
                )

            planilha.column_dimensions[letra].width = largura

    return buffer.getvalue()


def test_xlsx_datas_sem_horario():
    """
    Garante que as datas exportadas sejam legíveis
    e que os valores financeiros permaneçam iguais.
    """
    from datetime import datetime
    from openpyxl import load_workbook

    dados = pd.DataFrame(
        {
            "Dt. Emissão": pd.to_datetime(
                ["2026-06-12", "2026-07-13"]
            ),
            "Dt. Vencimento": pd.to_datetime(
                ["2026-09-15", "2026-09-30"]
            ),
            "Vlr Pago": [2356.00, 2028.00],
        }
    )

    original = dados.copy(deep=True)

    conteudo = exportar_xlsx(dados)
    planilha = load_workbook(
        io.BytesIO(conteudo)
    ).active

    # Datas preservadas sem horário significativo.
    assert planilha["A2"].value == datetime(2026, 6, 12)
    assert planilha["A3"].value == datetime(2026, 7, 13)

    assert planilha["A2"].number_format == "DD/MM/YYYY"
    assert planilha["B2"].number_format == "DD/MM/YYYY"

    # Colunas possuem largura suficiente.
    assert planilha.column_dimensions["A"].width >= 18
    assert planilha.column_dimensions["B"].width >= 18

    # Valores financeiros continuam numéricos.
    assert planilha["C2"].value == pytest.approx(2356.00)
    assert planilha["C3"].value == pytest.approx(2028.00)

    # A exportação não modifica os dados originais.
    pd.testing.assert_frame_equal(dados, original)
