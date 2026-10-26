
"""
Testes da exportação de dados tratados.

Não utiliza relatórios reais nem grava arquivos em disco.
"""

import io

import pandas as pd
import pytest

from src.core.exportacao import exportar_csv, exportar_xlsx


@pytest.fixture
def dados_sinteticos():
    """Cria dados financeiros sintéticos para os testes."""
    return pd.DataFrame(
        {
            "Pessoa": ["Fornecedor A", "Fornecedor B"],
            "Empresa": [1, 2],
            "Dt. Vencimento": pd.to_datetime(
                ["2026-10-26", "2026-10-30"]
            ),
            "Vlr Aberto": [1234.56, 250.75],
        }
    )


def test_csv_preserva_valores(dados_sinteticos):
    """Confere valores e quantidade de linhas após exportar CSV."""
    conteudo = exportar_csv(dados_sinteticos)

    resultado = pd.read_csv(
        io.BytesIO(conteudo),
        sep=";",
        decimal=",",
        encoding="utf-8-sig",
    )

    assert len(resultado) == 2
    assert resultado["Vlr Aberto"].sum() == pytest.approx(
        1485.31
    )
    assert resultado.loc[0, "Empresa"] == 1


def test_csv_utiliza_padrao_brasileiro(dados_sinteticos):
    """Verifica separador, decimal e codificação."""
    conteudo = exportar_csv(dados_sinteticos)
    texto = conteudo.decode("utf-8-sig")

    assert "1234,56" in texto
    assert "250,75" in texto
    assert "26/10/2026" in texto
    assert ";" in texto


def test_xlsx_preserva_valores(dados_sinteticos):
    """Confere os totais financeiros após gerar o Excel."""
    conteudo = exportar_xlsx(dados_sinteticos)

    resultado = pd.read_excel(
        io.BytesIO(conteudo),
        engine="openpyxl",
    )

    assert len(resultado) == 2
    assert resultado["Vlr Aberto"].sum() == pytest.approx(
        1485.31
    )


def test_xlsx_preserva_datas(dados_sinteticos):
    """Garante que datas continuam sendo datas no XLSX."""
    conteudo = exportar_xlsx(dados_sinteticos)

    resultado = pd.read_excel(
        io.BytesIO(conteudo),
        engine="openpyxl",
    )

    assert pd.api.types.is_datetime64_any_dtype(
        resultado["Dt. Vencimento"]
    )


def test_exportacao_nao_altera_original(dados_sinteticos):
    """A exportação nunca deve modificar a base filtrada."""
    original = dados_sinteticos.copy(deep=True)

    exportar_csv(dados_sinteticos)
    exportar_xlsx(dados_sinteticos)

    pd.testing.assert_frame_equal(
        dados_sinteticos, original
    )


def test_protecao_contra_formulas():
    """Textos provenientes do ERP não devem virar fórmulas."""
    dados = pd.DataFrame(
        {
            "Pessoa": ["=1+1", "+SUM(A1:A2)"],
            "Valor": [100.0, 200.0],
        }
    )

    csv = exportar_csv(dados).decode("utf-8-sig")
    xlsx = pd.read_excel(
        io.BytesIO(exportar_xlsx(dados)),
        engine="openpyxl",
    )

    assert "'=1+1" in csv
    assert "'+SUM(A1:A2)" in csv
    assert xlsx.loc[0, "Pessoa"] == "'=1+1"
    assert xlsx.loc[1, "Pessoa"] == "'+SUM(A1:A2)"


def test_xlsx_formato_celulas_datas():
    """
    Confere os valores e a formatação das datas no Excel.

    Não basta verificar se a data foi exportada:
    precisamos conferir como a célula foi configurada.
    """
    from datetime import datetime
    from openpyxl import load_workbook

    dados = pd.DataFrame(
        {
            "Dt. Emissão": pd.to_datetime(
                ["2026-09-01"]
            ),
            "Dt. Movimento": pd.to_datetime(
                ["2026-09-15"]
            ),
            "Dt. Quitação": pd.to_datetime(
                ["2026-09-22"]
            ),
            "Dt. Vencimento": pd.to_datetime(
                ["2026-09-30"]
            ),
            "Vlr Pago": [2356.00],
        }
    )

    conteudo = exportar_xlsx(dados)

    planilha = load_workbook(
        io.BytesIO(conteudo)
    ).active

    for coluna in range(1, 5):
        celula = planilha.cell(row=2, column=coluna)

        assert isinstance(celula.value, datetime)
        assert celula.number_format == "DD/MM/YYYY"

        letra = celula.column_letter

        assert (
            planilha.column_dimensions[letra].width >= 15
        )

    assert planilha["A2"].value == datetime(2026, 9, 1)
    assert planilha["D2"].value == datetime(2026, 9, 30)
    assert planilha["E2"].value == pytest.approx(2356.00)
