
"""
Diagnóstico local da formatação de arquivos XLSX.

Não imprime valores financeiros nem dados dos fornecedores.
Não modifica o arquivo examinado.
"""

from pathlib import Path

from openpyxl import load_workbook


def diagnosticar(caminho: str) -> None:
    """Exibe apenas metadados das colunas de datas."""
    arquivo = Path(caminho)

    planilha = load_workbook(
        arquivo, read_only=False, data_only=False
    ).active

    print("Diagnóstico das colunas de datas")
    print("-" * 45)

    for coluna in planilha.iter_cols():
        cabecalho = coluna[0].value

        if not isinstance(cabecalho, str):
            continue

        if not (
            cabecalho.startswith("Dt.")
            or "Vencimento" in cabecalho
        ):
            continue

        letra = coluna[0].column_letter
        largura = planilha.column_dimensions[letra].width

        print(f"\nColuna: {cabecalho}")
        print(f"Largura: {largura}")

        for celula in coluna[1:4]:
            print(
                f"Linha {celula.row}: "
                f"tipo={celula.data_type}, "
                f"formato={celula.number_format}, "
                f"estilo={celula.style_id}"
            )


if __name__ == "__main__":
    caminho = input("Caminho completo do XLSX: ").strip()
    diagnosticar(caminho)
